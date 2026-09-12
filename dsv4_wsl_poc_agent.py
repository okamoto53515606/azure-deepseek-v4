#!/usr/bin/env python3
"""Azure 上の DeepSeek V4 と BigQuery を順番に確認する PoC サンプル。

初めて実行するときは、次の順番で進める。

1. ``--ping-llm``: Azure AI Foundry だけに接続する（GCP の設定は不要）。
2. ``--check-config``: GCP 設定後に設定値と鍵ファイルを確認する。
3. ``--ping-bq``: BigQuery だけに ``SELECT 1`` を実行する。
4. ``--ping-bq-mcp``: BigQuery 用 MCP Toolbox の起動だけを確認する。
5. ``--bigquery-mcp``: AI エージェントから MCP 経由で BigQuery を使う。
6. オプションなし: AI エージェントから BigQuery SDK とレポート作成ツールを使う。
7. ``--sql-server-only``: GCP を使わず SQL Server とレポート機能だけを使う。

Azure Foundry には OpenAI 互換の Chat Completions で接続するため、Strands の
``OpenAIModel`` を使う。``OpenAIResponsesModel`` は使わない。

設定は実行プロセスの環境変数から直接読み、``config.py`` や ``.env`` は読み込まない。
BigQuery の認証にはチュートリアルで作成する JSON 鍵を使う。SQL Server は
MCP Toolbox の組み込み設定 ``--prebuilt mssql`` を使い、設定ファイルを作成しない。
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import shutil
import sys
from contextlib import ExitStack
from datetime import datetime, timedelta, timezone
from pathlib import Path, PureWindowsPath

from strands import Agent, tool
from strands.models.openai import OpenAIModel

AZURE_DIR = Path(__file__).resolve().parent
JST = timezone(timedelta(hours=9))
WORK_DIR = AZURE_DIR / "generated_files" / "dsv4_wsl_poc"
WINDOWS_C_MOUNT = Path("/mnt/c")
DEFAULT_WINDOWS_REPORT_DIR = WINDOWS_C_MOUNT / "DeepSeekV4PoC" / "reports"
DEFAULT_AZURE_AI_DEPLOYMENT = "DeepSeek-V4-Flash-0731"
DEFAULT_GCP_KEY_PATH = AZURE_DIR / "xxx.json"
FORBIDDEN_SQL = re.compile(
    r"\b(INSERT|UPDATE|DELETE|MERGE|DROP|ALTER|TRUNCATE|CREATE|GRANT|REVOKE|EXECUTE|CALL)\b",
    re.IGNORECASE,
)
DEFAULT_PROMPT = (
    "日本語で進めてください。"
    "1) 利用可能な BigQuery ツールで SELECT 1 を実行する。"
    "2) 利用可能な BigQuery ツールでデータセット名を最大8件取得する。"
    "3) write_work_report で作業レポート.docx を書く（表に ping 結果を入れる）。"
    "4) cleanup_work_files で一時ファイルを消し、docx のパスだけ残す。"
    "5) windows_explorer_path が返った場合は、最後の回答で Windows 用パスを案内する。"
    "APIキーや秘密は出力しない。SQL は読み取りのみ。"
)
BIGQUERY_MCP_TOOL_NAMES = {
    "execute_sql",
    "get_dataset_info",
    "get_table_info",
    "list_dataset_ids",
    "list_table_ids",
}
SQL_SERVER_ENV_NAMES = (
    "MSSQL_HOST",
    "MSSQL_PORT",
    "MSSQL_DATABASE",
    "MSSQL_USER",
    "MSSQL_PASSWORD",
)
SQL_SERVER_MCP_TOOL_NAMES = {"execute_sql", "list_tables"}

_ASYNCIO_SHUTDOWN_NOISE = "an error occurred during closing of asynchronous generator"


class _AsyncGeneratorShutdownNoiseFilter(logging.Filter):
    """終了時に出る既知のノイズを非表示にする。

    Strands の OpenAI モデルは応答ストリームを最後まで読み切らないため、
    イベントループの終了時に httpcore2 の非同期ジェネレーターが閉じられる。
    そのとき asyncio が「an error occurred during closing of asynchronous
    generator ...」とトレースバックをログへ出力するが、回答や処理結果には
    影響しない。このメッセージ 1 件だけを除外する。
    """

    def filter(self, record: logging.LogRecord) -> bool:
        return _ASYNCIO_SHUTDOWN_NOISE not in record.getMessage()


def _suppress_asyncgen_shutdown_noise() -> None:
    """asyncio ロガーへ上記フィルターを 1 度だけ取り付ける。"""
    logger = logging.getLogger("asyncio")
    if not any(isinstance(item, _AsyncGeneratorShutdownNoiseFilter) for item in logger.filters):
        logger.addFilter(_AsyncGeneratorShutdownNoiseFilter())


def azure_ai_endpoint() -> str:
    return (os.environ.get("AZURE_AI_ENDPOINT") or "").strip().rstrip("/")


def azure_ai_api_key() -> str:
    key = (os.environ.get("AZURE_AI_API_KEY") or "").strip()
    if not key:
        raise ValueError("AZURE_AI_API_KEY が未設定。Azure の API キーを設定してください")
    return key


def azure_chat_deployment() -> str:
    return (os.environ.get("AZURE_AI_DEPLOYMENT") or DEFAULT_AZURE_AI_DEPLOYMENT).strip()


def azure_openai_base_url() -> str:
    return f"{azure_ai_endpoint()}/openai/v1/"


def gcp_credentials_path() -> str | None:
    configured_path = (os.environ.get("GOOGLE_APPLICATION_CREDENTIALS") or "").strip()
    if configured_path:
        path = Path(configured_path).expanduser()
        return str(path) if path.is_file() else None
    return str(DEFAULT_GCP_KEY_PATH) if DEFAULT_GCP_KEY_PATH.is_file() else None


def _validate_azure_settings() -> tuple[str, str]:
    """Azure 接続に必要な値を確認する。API キーの値自体は返さない。"""
    endpoint = azure_ai_endpoint()
    if not endpoint:
        raise ValueError("AZURE_AI_ENDPOINT が未設定。Azure のエンドポイントを設定してください")
    if not endpoint.startswith("https://"):
        raise ValueError("AZURE_AI_ENDPOINT は https:// から始まる URL を設定してください")

    deployment = azure_chat_deployment().strip()
    if not deployment:
        raise ValueError("AZURE_AI_DEPLOYMENT が未設定。Azure のデプロイ名を設定してください")

    azure_ai_api_key()
    return endpoint, deployment


def _setup_gcp_sa() -> str:
    """ローカルの SA JSON を GOOGLE_APPLICATION_CREDENTIALS にする。中身（秘密鍵）は出さない。"""
    path = gcp_credentials_path()
    if not path:
        raise FileNotFoundError(
            "GCP サービスアカウントの JSON 鍵が見つかりません。"
            " azure/xxx.json に置くか、"
            " GOOGLE_APPLICATION_CREDENTIALS にファイルのパスを設定してください"
        )
    os.environ["GOOGLE_APPLICATION_CREDENTIALS"] = path
    return path


def _sa_public_info(path: str) -> tuple[str, str]:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"GCP の JSON 鍵を読み取れません: {path}") from exc
    if data.get("type") != "service_account":
        raise ValueError("GCP の JSON は type=service_account の鍵を指定してください")
    project = (os.environ.get("GOOGLE_BIGQUERY_PROJECT") or data.get("project_id") or "").strip()
    email = (data.get("client_email") or "").strip()
    if not project:
        raise ValueError("GCP の JSON 鍵に project_id がありません")
    if not email:
        raise ValueError("GCP の JSON 鍵に client_email がありません")
    return project, email


def _bq_client():
    try:
        from google.cloud import bigquery
    except ImportError as exc:
        raise RuntimeError(
            "google-cloud-bigquery が未インストールです。"
            " 手順書の uv run --with google-cloud-bigquery ... で実行してください"
        ) from exc

    path = os.environ["GOOGLE_APPLICATION_CREDENTIALS"]
    project, _email = _sa_public_info(path)
    location = os.environ.get("GOOGLE_BIGQUERY_LOCATION", "asia-northeast1")
    return bigquery.Client.from_service_account_json(path, project=project), location


def _gcs_client():
    try:
        from google.cloud import storage
    except ImportError as exc:
        raise RuntimeError(
            "google-cloud-storage が未インストールです。"
            " 手順書の uv run --with google-cloud-storage ... で実行してください"
        ) from exc

    path = os.environ["GOOGLE_APPLICATION_CREDENTIALS"]
    project, _email = _sa_public_info(path)
    return storage.Client.from_service_account_json(path, project=project)


def _assert_select_only(sql: str) -> str:
    text = (sql or "").strip().rstrip(";")
    if not text:
        raise ValueError("SQL が空")
    if FORBIDDEN_SQL.search(text):
        raise ValueError("読み取り専用。DML/DDL は禁止")
    if not re.match(r"^(WITH|SELECT)\b", text, re.IGNORECASE | re.DOTALL):
        raise ValueError("先頭は SELECT または WITH のみ")
    return text


@tool(name="bq_ping", description="BigQuery に SELECT 1 を投げて接続確認する。引数なし。")
def bq_ping() -> str:
    client, location = _bq_client()
    job = client.query("SELECT 1 AS ok", location=location)
    rows = list(job.result(timeout=60))
    ok = int(rows[0]["ok"]) if rows else 0
    return json.dumps(
        {"ok": ok, "project": client.project, "location": location},
        ensure_ascii=False,
    )


@tool(name="bq_list_datasets", description="プロジェクト内データセット ID を最大 limit 件返す。")
def bq_list_datasets(limit: int = 8) -> str:
    client, _location = _bq_client()
    n = max(1, min(int(limit), 50))
    ids = []
    for ds in client.list_datasets():
        ids.append(ds.dataset_id)
        if len(ids) >= n:
            break
    return json.dumps({"project": client.project, "datasets": ids}, ensure_ascii=False)


@tool(
    name="bq_run_select",
    description="BigQuery で SELECT（または WITH）だけ実行する。最大 max_rows 行。DML/DDL 禁止。",
)
def bq_run_select(sql: str, max_rows: int = 50) -> str:
    text = _assert_select_only(sql)
    client, location = _bq_client()
    n = max(1, min(int(max_rows), 200))
    job = client.query(text, location=location)
    result = job.result(timeout=120)
    cols = [f.name for f in result.schema] if result.schema else []
    rows = []
    truncated = False
    for index, row in enumerate(result):
        if index >= n:
            truncated = True
            break
        rows.append({c: row[c] for c in cols})
    return json.dumps(
        {"project": client.project, "location": location, "columns": cols, "rows": rows, "truncated": truncated},
        ensure_ascii=False,
        default=str,
    )


@tool(name="gcs_list_blobs", description="GCS バケットのオブジェクト名を prefix 付きで最大 limit 件。読み取りのみ。")
def gcs_list_blobs(bucket: str, prefix: str = "", limit: int = 20) -> str:
    name = (bucket or "").strip()
    if not name:
        raise ValueError("bucket が空")
    n = max(1, min(int(limit), 100))
    client = _gcs_client()
    names = []
    for blob in client.list_blobs(name, prefix=prefix or None, max_results=n):
        names.append(blob.name)
    return json.dumps({"bucket": name, "prefix": prefix, "blobs": names}, ensure_ascii=False)


@tool(
    name="gcs_download_text",
    description="GCS のテキストオブジェクトを作業ディレクトリへ保存し、先頭を返す。最大 1MB。",
)
def gcs_download_text(bucket: str, blob_name: str) -> str:
    name = (bucket or "").strip()
    key = (blob_name or "").strip().lstrip("/")
    if not name or not key or ".." in key:
        raise ValueError("bucket / blob_name が不正")
    client = _gcs_client()
    blob = client.bucket(name).blob(key)
    data = blob.download_as_bytes()
    if len(data) > 1_000_000:
        raise ValueError("1MB 超は扱わない")
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    dest = WORK_DIR / Path(key).name
    dest.write_bytes(data)
    preview = data[:2000].decode("utf-8", errors="replace")
    return json.dumps(
        {"saved": str(dest), "bytes": len(data), "preview": preview},
        ensure_ascii=False,
    )


@tool(
    name="write_chart_png",
    description="labels と values から棒グラフ PNG を作業ディレクトリに書く。日本語ラベルは文字化けし得る。",
)
def write_chart_png(labels: list[str], values: list[float], filename: str = "chart.png") -> str:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    safe = Path(filename).name
    if not safe.endswith(".png"):
        safe += ".png"
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    path = WORK_DIR / safe
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.bar(list(map(str, labels)), [float(v) for v in values])
    ax.set_title("poc chart")
    fig.tight_layout()
    fig.savefig(path, dpi=120)
    plt.close(fig)
    return str(path)


@tool(
    name="write_work_report",
    description=(
        "Word レポートを書く。paragraphs は本文段落、table_rows は先頭行をヘッダにした表、"
        "data_source はレポートに表示するデータソース名。WSL の保存先と、利用可能なら "
        "Windows の C ドライブへコピーした保存先を返す。"
    ),
)
def write_work_report(
    title: str,
    paragraphs: list[str],
    table_rows: list[list[str]] | None = None,
    filename: str = "作業レポート.docx",
    data_source: str = "BigQuery",
    copy_to_windows: bool = True,
) -> str:
    from docx import Document

    safe = Path(filename).name
    if not safe.endswith(".docx"):
        safe += ".docx"
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    path = WORK_DIR / safe
    doc = Document()
    doc.add_heading(title or "作業レポート", level=1)
    now = datetime.now(tz=JST).strftime("%Y-%m-%d %H:%M:%S JST")
    doc.add_paragraph(f"作成: {now}")
    doc.add_paragraph("モデル: Azure Foundry DeepSeek-V4-Flash-0731（OpenAIModel / Chat Completions）")
    doc.add_paragraph(f"データソース: {data_source or '未指定'}")
    for p in paragraphs or []:
        if p:
            doc.add_paragraph(str(p))
    if table_rows:
        cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=cols)
        table.style = "Table Grid"
        for i, row in enumerate(table_rows):
            for j in range(cols):
                table.cell(i, j).text = str(row[j]) if j < len(row) else ""
    chart = WORK_DIR / "chart.png"
    if chart.is_file():
        from docx.shared import Inches

        doc.add_paragraph("グラフ:")
        doc.add_picture(str(chart), width=Inches(5.5))
    doc.save(path)

    result = {"wsl_path": str(path)}
    if copy_to_windows and WINDOWS_C_MOUNT.is_dir():
        windows_dir = Path(
            os.environ.get("WINDOWS_REPORT_DIR", str(DEFAULT_WINDOWS_REPORT_DIR))
        ).expanduser().resolve()
        try:
            windows_dir.relative_to(WINDOWS_C_MOUNT.resolve())
        except ValueError as exc:
            raise ValueError("WINDOWS_REPORT_DIR は /mnt/c/ 配下を指定してください") from exc
        windows_dir.mkdir(parents=True, exist_ok=True)
        windows_copy = windows_dir / safe
        shutil.copy2(path, windows_copy)
        relative = windows_copy.relative_to(WINDOWS_C_MOUNT)
        result["windows_path"] = str(windows_copy)
        result["windows_explorer_path"] = str(PureWindowsPath("C:/", *relative.parts))
    return json.dumps(result, ensure_ascii=False)


@tool(
    name="cleanup_work_files",
    description="作業ディレクトリの一時ファイルを消す。keep_docx=true なら .docx は残す。",
)
def cleanup_work_files(keep_docx: bool = True) -> str:
    if not WORK_DIR.is_dir():
        return json.dumps({"deleted": [], "kept": []}, ensure_ascii=False)
    deleted = []
    kept = []
    for p in WORK_DIR.iterdir():
        if not p.is_file():
            continue
        if keep_docx and p.suffix.lower() == ".docx":
            kept.append(str(p))
            continue
        p.unlink()
        deleted.append(p.name)
    return json.dumps({"work_dir": str(WORK_DIR), "deleted": deleted, "kept": kept}, ensure_ascii=False)


def _build_model() -> OpenAIModel:
    return OpenAIModel(
        client_args={
            "api_key": azure_ai_api_key(),
            "base_url": azure_openai_base_url(),
            "timeout": 300.0,
            "max_retries": 0,
        },
        model_id=azure_chat_deployment(),
        params={"max_tokens": 8192},
    )


def _toolbox_command() -> str:
    """Toolbox の実行ファイルがローカルにあることを確認する。"""
    command = shutil.which("toolbox")
    if not command:
        raise FileNotFoundError(
            "toolbox コマンドが見つかりません。MCP Toolbox をインストールし、PATH を設定してください"
        )
    return command


def _validate_sql_server_settings() -> None:
    missing = [name for name in SQL_SERVER_ENV_NAMES if not (os.environ.get(name) or "").strip()]
    if missing:
        raise ValueError(f"SQL Server の環境変数が未設定です: {', '.join(missing)}")
    try:
        port = int(os.environ["MSSQL_PORT"])
    except ValueError as exc:
        raise ValueError("MSSQL_PORT は数字で設定してください") from exc
    if not 1 <= port <= 65535:
        raise ValueError("MSSQL_PORT は 1〜65535 の範囲で設定してください")


def _build_toolbox_mcp():
    """Toolbox の SQL Server 組み込み設定を stdio で起動する。"""
    from mcp import StdioServerParameters, stdio_client
    from strands.tools.mcp import MCPClient

    _validate_sql_server_settings()
    command = _toolbox_command()
    return MCPClient(
        lambda: stdio_client(
            StdioServerParameters(
                command=command,
                args=["--prebuilt", "mssql", "--stdio"],
                env=os.environ.copy(),
            )
        )
    )


def _build_bigquery_mcp():
    """Toolbox の BigQuery 組み込み設定を stdio MCP として起動する。"""
    from mcp import StdioServerParameters, stdio_client
    from strands.tools.mcp import MCPClient

    command = _toolbox_command()

    path = os.environ["GOOGLE_APPLICATION_CREDENTIALS"]
    project, _email = _sa_public_info(path)
    env = os.environ.copy()
    env["BIGQUERY_PROJECT"] = project
    return MCPClient(
        lambda: stdio_client(
            StdioServerParameters(
                command=command,
                args=["--prebuilt", "bigquery", "--stdio"],
                env=env,
            )
        )
    )


def ping_toolbox() -> int:
    """Toolbox を起動し、MCP ツール一覧を取得できることを確認する。DB へ SQL は送らない。"""
    toolbox_mcp = _build_toolbox_mcp()
    with toolbox_mcp:
        database_tools = toolbox_mcp.list_tools_sync()
    tool_names = {tool.tool_name for tool in database_tools}
    missing = SQL_SERVER_MCP_TOOL_NAMES.difference(tool_names)
    if missing:
        raise RuntimeError(f"SQL Server MCP に必要なツールがありません: {', '.join(sorted(missing))}")
    return len(SQL_SERVER_MCP_TOOL_NAMES)


def ping_bigquery_mcp() -> list[str]:
    """BigQuery MCP を起動し、サンプルで許可するツール名を返す。"""
    bigquery_mcp = _build_bigquery_mcp()
    with bigquery_mcp:
        names = sorted(
            tool.tool_name
            for tool in bigquery_mcp.list_tools_sync()
            if tool.tool_name in BIGQUERY_MCP_TOOL_NAMES
        )
    missing = BIGQUERY_MCP_TOOL_NAMES.difference(names)
    if missing:
        raise RuntimeError(f"BigQuery MCP に必要なツールがありません: {', '.join(sorted(missing))}")
    return names


SYSTEM_PROMPT = """あなたは検証用の分析アシスタント。
- メインモデルは Azure AI Foundry の DeepSeek-V4-Flash-0731。Strands は OpenAIModel（Chat Completions）。
- SQL は読み取りのみ。秘密・APIキー・SA JSON の中身は出力しない。
- MCP Toolbox の SQL ツールがある場合も SELECT / WITH だけを使う。
- レポートを書いたら cleanup_work_files で一時ファイルを消す（docx は残してよい）。
- ツールが失敗したら推測で埋めず、エラー内容を短く返す。
"""


def ping_llm() -> str:
    from openai import OpenAI

    client = OpenAI(
        api_key=azure_ai_api_key(),
        base_url=azure_openai_base_url(),
        timeout=60.0,
    )
    r = client.chat.completions.create(
        model=azure_chat_deployment(),
        messages=[{"role": "user", "content": "1+1は？数字だけ答えて。"}],
        max_tokens=32,
    )
    return (r.choices[0].message.content or "").strip()


def _core_tools(include_bigquery_sdk: bool = True, include_gcs: bool = True) -> list:
    tools = [
        write_chart_png,
        write_work_report,
        cleanup_work_files,
    ]
    if include_gcs:
        tools[:0] = [gcs_list_blobs, gcs_download_text]
    if include_bigquery_sdk:
        tools[:0] = [bq_ping, bq_list_datasets, bq_run_select]
    return tools


def _run_agent(
    prompt: str,
    sql_server: bool = False,
    bigquery_mcp: bool = False,
    include_bigquery_sdk: bool = True,
    include_gcs: bool = True,
):
    tools = _core_tools(
        include_bigquery_sdk=include_bigquery_sdk and not bigquery_mcp,
        include_gcs=include_gcs,
    )
    with ExitStack() as stack:
        if bigquery_mcp:
            bigquery_client = stack.enter_context(_build_bigquery_mcp())
            bigquery_tools = [
                tool
                for tool in bigquery_client.list_tools_sync()
                if tool.tool_name in BIGQUERY_MCP_TOOL_NAMES
            ]
            missing = BIGQUERY_MCP_TOOL_NAMES.difference(tool.tool_name for tool in bigquery_tools)
            if missing:
                raise RuntimeError(f"BigQuery MCP に必要なツールがありません: {', '.join(sorted(missing))}")
            print(f"bigquery_mcp_tools={len(bigquery_tools)}", flush=True)
            tools.extend(bigquery_tools)

        if sql_server:
            toolbox_client = stack.enter_context(_build_toolbox_mcp())
            database_tools = [
                tool
                for tool in toolbox_client.list_tools_sync()
                if tool.tool_name in SQL_SERVER_MCP_TOOL_NAMES
            ]
            missing = SQL_SERVER_MCP_TOOL_NAMES.difference(tool.tool_name for tool in database_tools)
            if missing:
                raise RuntimeError(f"SQL Server MCP に必要なツールがありません: {', '.join(sorted(missing))}")
            print(f"toolbox_tools={len(database_tools)}", flush=True)
            tools.extend(database_tools)

        return Agent(model=_build_model(), tools=tools, system_prompt=SYSTEM_PROMPT)(prompt)


def _print_gcp_settings() -> tuple[str, str, str]:
    sa_path = _setup_gcp_sa()
    project, email = _sa_public_info(sa_path)
    location = os.environ.get("GOOGLE_BIGQUERY_LOCATION", "asia-northeast1")
    print(f"gcp_sa_json={sa_path}", flush=True)
    print(f"gcp_project={project}", flush=True)
    print(f"gcp_sa_email={email}", flush=True)
    print(f"gcp_location={location}", flush=True)
    return sa_path, project, email


def _troubleshooting_hint(error: Exception) -> str:
    message = str(error).lower()
    if "401" in message or "authentication" in message or "api key" in message or "api_key" in message:
        return "確認: AZURE_AI_API_KEY が正しいか、期限切れやコピー漏れがないか確認してください。"
    if "404" in message:
        return "確認: AZURE_AI_ENDPOINT と AZURE_AI_DEPLOYMENT が同じ Azure リソースの値か確認してください。"
    if "permission" in message or "403" in message:
        return "確認: 実行ユーザーまたはサービスアカウントに必要な閲覧権限があるか確認してください。"
    if "name resolution" in message or "connection" in message or "timeout" in message:
        return "確認: 会社プロキシ、VPN、DNS、接続先ホスト名とポートを確認してください。"
    return "確認: 手順書のトラブルシューティング表と、上に表示された設定名を確認してください。"


def main() -> int:
    _suppress_asyncgen_shutdown_noise()
    parser = argparse.ArgumentParser(
        description="Azure DeepSeek V4 + BigQuery の段階確認用サンプル",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""実行順:
    1. --ping-llm  2. --check-config  3. --ping-bq  4. --ping-bq-mcp
    5. --bigquery-mcp（DeepSeek から MCP 経由で BigQuery を使用）
SQL Server は --ping-toolbox で確認してから、通常実行に --sql-server-only を追加します。""",
    )
    parser.add_argument("--prompt", default=DEFAULT_PROMPT)
    check_group = parser.add_mutually_exclusive_group()
    check_group.add_argument(
        "--check-config",
        action="store_true",
        help="設定値を確認（Azure AI / BigQuery へのテスト要求なし）",
    )
    check_group.add_argument("--ping-llm", action="store_true", help="Azure AI にだけ接続して 1+1 を確認")
    check_group.add_argument("--ping-bq", action="store_true", help="BigQuery にだけ SELECT 1 を実行")
    check_group.add_argument(
        "--ping-toolbox",
        action="store_true",
        help="Toolbox を起動して MCP ツール一覧を取得（SQL は実行しない）",
    )
    check_group.add_argument(
        "--ping-bq-mcp",
        action="store_true",
        help="BigQuery 用 Toolbox を起動して MCP ツール一覧を取得（SQL は実行しない）",
    )
    parser.add_argument(
        "--bigquery-mcp",
        action="store_true",
        help="通常実行で BigQuery を直接 SDK ではなく MCP Toolbox 経由にする",
    )
    parser.add_argument(
        "--sql-server-only",
        action="store_true",
        help="通常実行で GCP を使わず SQL Server の Toolbox とレポート機能だけを使う",
    )
    args = parser.parse_args()

    if args.ping_toolbox:
        tool_count = ping_toolbox()
        print(f"[OK] Toolbox が起動しました: tools={tool_count}", flush=True)
        print("[INFO] この確認では SQL Server への SQL 実行は行っていません。", flush=True)
        return 0

    if args.ping_bq_mcp:
        _print_gcp_settings()
        tool_names = ping_bigquery_mcp()
        print(f"[OK] BigQuery MCP が起動しました: tools={','.join(tool_names)}", flush=True)
        print("[INFO] この確認では BigQuery への SQL 実行は行っていません。", flush=True)
        return 0

    # BigQuery だけを確認する段階では、Azure や Key Vault への接続を行わない。
    if args.ping_bq:
        _print_gcp_settings()
        print(f"bq_response={bq_ping()}")
        return 0

    if args.sql_server_only and args.bigquery_mcp:
        parser.error("--sql-server-only と --bigquery-mcp は同時に指定できません")

    if args.check_config:
        endpoint, deployment = _validate_azure_settings()
        print(f"[OK] Azure AI: endpoint={endpoint} deployment={deployment}", flush=True)
        if not args.sql_server_only:
            _sa_path, project, email = _print_gcp_settings()
            print(f"[OK] GCP: project={project} service_account={email}", flush=True)
        if args.sql_server_only:
            _validate_sql_server_settings()
            print(f"[OK] Toolbox: command={_toolbox_command()} config=--prebuilt mssql", flush=True)
        print("[OK] 設定を確認しました。Azure AI と BigQuery へのテスト要求はまだ送っていません。", flush=True)
        return 0

    # Azure だけを確認する段階では、GCP の設定や鍵を要求しない。
    if args.ping_llm:
        endpoint, deployment = _validate_azure_settings()
        print(f"azure_endpoint={endpoint} deployment={deployment}", flush=True)
        print(f"llm_response={ping_llm()}")
        return 0

    endpoint, deployment = _validate_azure_settings()
    if not args.sql_server_only:
        _print_gcp_settings()
    print(f"azure_endpoint={endpoint} deployment={deployment}", flush=True)
    print(f"work_dir={WORK_DIR}", flush=True)

    WORK_DIR.mkdir(parents=True, exist_ok=True)
    result = _run_agent(
        args.prompt,
        sql_server=args.sql_server_only,
        bigquery_mcp=args.bigquery_mcp,
        include_bigquery_sdk=not args.sql_server_only,
        include_gcs=not args.sql_server_only,
    )
    # 最終回答はコールバックハンドラーがストリーミング表示済み。ここで print すると
    # 同じ内容が 2 回表示されるため、結果の有無だけ確認する。
    if result is None:
        print("[WARN] エージェントから結果が返りませんでした", file=sys.stderr)
    # ストリーミング表示は改行で終わらないことがあるため、シェルのプロンプトと
    # 重ならないように改行を補う。
    print(flush=True)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"エラー: {type(error).__name__}: {error}", file=sys.stderr)
        print(_troubleshooting_hint(error), file=sys.stderr)
        raise SystemExit(1)
